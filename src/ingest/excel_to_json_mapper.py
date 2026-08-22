#!/usr/bin/env python3
"""Ingest the FIT AI exercise toolkit + squat CSV into a tagged catalog."""

from __future__ import annotations

import csv
import json
import os
import re
from pathlib import Path

from openpyxl import load_workbook

from src.taxonomy import (
    CATEGORY_ALIASES,
    SHEET_PATTERN,
    SHEET_PROGRAM_ROLE,
    TYPO_FIXES,
    generate_smart_tags,
)

ROOT = Path(__file__).resolve().parents[2]
INPUT_XLSX = ROOT / "data/raw/FIT_AI_Exercise_Description_Toolkit.xlsx"
INPUT_CSV = ROOT / "data/raw/squat_data_sheet.csv"
OUTPUT_JSON = ROOT / "data/processed/exercise_catalog.json"
LEGACY_JSON = ROOT / "data/processed/exercise_knowledge_base.json"

LEVEL_RE = re.compile(r"Level\s+(\d+)", re.I)
BOILERPLATE_RE = re.compile(
    r"This variation manipulates leverage, loading, stability", re.I
)

ACRONYM_EQUIPMENT = {
    "BW": "bodyweight",
    "DB": "dumbbell",
    "KB": "kettlebell",
    "BB": "barbell",
    "TRX": "trx",
    "MB": "medball",
    "SB": "sandbag",
    "LM": "landmine",
    "SM": "smith",
    "BAND": "band",
    "BANDED": "band",
    "WALL": "wall",
    "CABLE": "cable",
    "PLATE": "plate",
}


def slug(text: str) -> str:
    text = (text or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_") or "unnamed"


def normalize_typos(text: str) -> str:
    out = text or ""
    for bad, good in TYPO_FIXES.items():
        out = re.sub(re.escape(bad), good, out, flags=re.I)
    return out


def parse_ramp_role(family: str) -> str | None:
    upper = family.upper()
    if "RAISE" in upper:
        return "raise"
    if "ACTIVATE" in upper:
        return "activate"
    if "MOBIL" in upper:
        return "mobilize"
    if "POTENTIATE" in upper:
        return "potentiate"
    return None


def parse_equipment(name: str) -> list[str]:
    tokens = set(re.findall(r"[A-Z]{2,}", name.upper()))
    found = []
    for token, eq in ACRONYM_EQUIPMENT.items():
        if token in tokens or token in name.upper():
            found.append(eq)
    if not found:
        found.append("bodyweight")
    return sorted(set(found))


def parse_laterality(name: str) -> str | None:
    upper = name.upper()
    if "U/L" in upper or "UNILATERAL" in upper or "SINGLE LEG" in upper or re.search(r"\bSL\b", upper):
        return "unilateral"
    if "B/L" in upper or "BILATERAL" in upper:
        return "bilateral"
    return None


def parse_muscles(text: str) -> dict:
    if not text:
        return {"primary": [], "secondary": []}
    primary, secondary = [], []
    m = re.search(r"Primary:\s*([^;]+)", text, re.I)
    if m:
        primary = [p.strip() for p in m.group(1).split(",") if p.strip()]
    m = re.search(r"Secondary:\s*([^;]+)", text, re.I)
    if m:
        secondary = [p.strip() for p in m.group(1).split(",") if p.strip()]
    return {"primary": primary, "secondary": secondary}


def is_boilerplate(desc: str) -> bool:
    return bool(desc and BOILERPLATE_RE.search(desc))


def load_squat_csv() -> dict[tuple[str, str], dict]:
    """Parse the messy squat CSV, repairing unquoted commas inside parentheses."""
    lookup = {}
    if not INPUT_CSV.exists():
        return lookup

    raw_rows = []
    with open(INPUT_CSV, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        header = next(reader, None)
        for row in reader:
            raw_rows.append(row)

    for row in raw_rows:
        if not row or not any(c.strip() for c in row):
            continue
        cells = [c.strip() for c in row]
        # Skip acronym footer
        if cells[0].upper() in {"B/L", "U/L", "BW", "DB", "KB", "BB"} or (
            not cells[0].isdigit() and not cells[1]
        ):
            if not (len(cells) > 1 and cells[1].upper().startswith("LEVEL")):
                continue

        # Reconstruct: SL NO, Category, Level, Variation, Description, Muscles, Notes
        if len(cells) < 4:
            continue
        sl, category, level, *rest = cells
        if not category or not str(level).upper().startswith("LEVEL"):
            continue

        variation = rest[0] if rest else ""
        # Merge split variation names with unclosed parentheses
        idx = 1
        while idx < len(rest) and variation.count("(") > variation.count(")"):
            variation = variation + "," + rest[idx]
            idx += 1
        leftover = rest[idx:]
        description = leftover[0] if leftover else ""
        muscles = leftover[1] if len(leftover) > 1 else ""
        notes = leftover[2] if len(leftover) > 2 else ""

        # If muscles field doesn't look like muscles, it may be notes
        if muscles and "Primary" not in muscles and not notes:
            notes, muscles = muscles, ""
            for item in leftover[1:]:
                if "Primary" in item:
                    muscles = item
                    notes = " ".join(x for x in leftover[1:] if x != item)
                    break

        key_cat = CATEGORY_ALIASES.get(category.upper(), category.upper())
        key_name = normalize_typos(variation).upper()
        lookup[(key_cat, key_name)] = {
            "category": category,
            "level": level,
            "variation": variation,
            "description": description,
            "muscles": parse_muscles(muscles),
            "indications_text": notes.strip(),
        }
    return lookup


def ingest_workbook() -> list[dict]:
    wb = load_workbook(INPUT_XLSX, data_only=True)
    squat_lookup = load_squat_csv()
    catalog = []
    seen_ids = set()

    for sheet_name, pattern in SHEET_PATTERN.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        family = "UNCATEGORIZED"
        program_role = SHEET_PROGRAM_ROLE.get(sheet_name, "accessory")

        for row in ws.iter_rows(min_row=2, values_only=True):
            name = normalize_typos(str(row[0]).strip()) if row[0] else ""
            desc = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            if not name or name.lower() == "exercise name":
                continue
            if name and not desc:
                family = name
                continue

            level_match = LEVEL_RE.search(desc)
            level = int(level_match.group(1)) if level_match else 1
            ramp_role = parse_ramp_role(family) if sheet_name == "WARM-UP" else None
            role = program_role
            if ramp_role == "activate":
                role = "activation"
            elif ramp_role in ("raise", "mobilize", "potentiate"):
                role = "warmup"

            csv_key = (
                CATEGORY_ALIASES.get(family.upper(), family.upper()),
                name.upper(),
            )
            csv_row = squat_lookup.get(csv_key)
            muscles = {"primary": [], "secondary": []}
            indications = ""
            if csv_row:
                muscles = csv_row["muscles"]
                indications = csv_row["indications_text"]
                if is_boilerplate(desc) and csv_row["description"]:
                    desc = csv_row["description"]

            exercise_id = slug(f"{sheet_name}_{family}_{name}_{level}")
            n = 2
            base_id = exercise_id
            while exercise_id in seen_ids:
                exercise_id = f"{base_id}_{n}"
                n += 1
            seen_ids.add(exercise_id)

            entry = {
                "id": exercise_id,
                "name": name,
                "family": family,
                "pattern": pattern,
                "aliases": [csv_row["variation"]] if csv_row and csv_row["variation"] != name else [],
                "ramp_role": ramp_role,
                "program_role": role,
                "level": level,
                "equipment": parse_equipment(name + " " + family),
                "laterality": parse_laterality(name),
                "description": desc,
                "muscles": muscles,
                "indications_text": indications,
                "source_sheet": sheet_name,
                "tags": generate_smart_tags(name, family, pattern, level, ramp_role),
                "impact": None,
                "load_basis": None,
                "contraindications": [],
                "regression_ids": [],
            }
            catalog.append(entry)

    return catalog


def run_ingestion():
    if not INPUT_XLSX.exists():
        print(f"❌ Missing toolkit: {INPUT_XLSX}")
        return
    os.makedirs(OUTPUT_JSON.parent, exist_ok=True)
    catalog = ingest_workbook()
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(catalog, f, indent=2)
    # Keep legacy path so old retriever still has data during transition
    with open(LEGACY_JSON, "w", encoding="utf-8") as f:
        json.dump(
            [
                {
                    "id": e["id"],
                    "exercise_name": e["name"],
                    "category": e["family"],
                    "difficulty_level": e["level"],
                    "description": e["description"],
                    "description_source": "ingested",
                    "tags": e["tags"],
                }
                for e in catalog
            ],
            f,
            indent=2,
        )
    print(f"✅ Ingested {len(catalog)} exercises → {OUTPUT_JSON}")


if __name__ == "__main__":
    run_ingestion()
